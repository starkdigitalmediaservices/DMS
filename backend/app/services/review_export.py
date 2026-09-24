"""Export a document's CORRECTED version from the review screen.

Built from exactly what the review screen shows (get_review_document), so an
export never disagrees with the screen. Removed rows/blocks are left out.
Every value carries its status, so a reader can tell human-checked data from
machine output:

  Checked by a person     the row/block was marked checked, or the value was
                          confirmed (Workbench "Correct")
  Corrected by a person   a reviewer changed it, not yet checked
  Added by a person       text/rows the OCR missed, typed in by a reviewer
  Machine-extracted       as the computer read it; nobody has looked yet

Formats: xlsx (a sheet per table, values next to their status, colour
coded, plus an all-values sheet), csv (one line per value -- works for any
document shape), json.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import quote
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

CHECKED = "Checked by a person"
CORRECTED = "Corrected by a person"
ADDED = "Added by a person"
MACHINE = "Machine-extracted"

_FILL = {CHECKED: "D9F2DF", CORRECTED: "FFF1C2", ADDED: "FFE3CC", MACHINE: None}
FORMATS = ("xlsx", "csv", "json")


def _label(header: str) -> str:
    """Machine field names -> readable (same rule as the screen); real
    headings read off the scan are kept exactly."""
    if not header or not re.fullmatch(r"[a-z0-9_]+", header):
        return header
    words = [f"(column {w[3:]})" if re.fullmatch(r"col\d+", w) else
             {"sr": "serial", "no": "no.", "wakf": "waqf", "mutawalli": "mutawalli (manager)"}.get(w, w)
             for w in header.split("_") if w]
    text = " ".join(words)
    return text[:1].upper() + text[1:]


def _cell_status(cell: Dict[str, Any], row: Dict[str, Any]) -> str:
    if row.get("status") == "VERIFIED" or cell.get("status") == "VERIFIED":
        return CHECKED
    if row.get("added"):
        return ADDED
    if cell.get("edited"):
        return CORRECTED
    return MACHINE


def _block_status(block: Dict[str, Any]) -> str:
    if block.get("status") == "VERIFIED":
        return CHECKED
    if block.get("added"):
        return ADDED
    if block.get("edited"):
        return CORRECTED
    return MACHINE


def _page(block: Dict[str, Any], row: Optional[Dict[str, Any]] = None) -> Optional[int]:
    if row and row.get("page"):
        return row["page"]
    return (block.get("source_pages") or [None])[0]


def flatten(review_doc: Dict[str, Any], names: Dict[str, str]) -> List[Dict[str, Any]]:
    """One record per value: every table cell and every text block."""
    out: List[Dict[str, Any]] = []
    table_no = 0
    for block in review_doc["blocks"]:
        if block.get("deleted"):
            continue
        if block["type"] == "table":
            table_no += 1
            headers = block.get("headers") or []
            live = [r for r in block.get("rows") or [] if not r.get("deleted")]
            for n, row in enumerate(live, start=1):
                for col, cell in enumerate(row["cells"]):
                    out.append({
                        "section": block.get("title") or f"Table {table_no}",
                        "section_type": "table",
                        "page": _page(block, row),
                        "row": n,
                        "field": _label(headers[col]) if col < len(headers) else f"Column {col + 1}",
                        "value": cell.get("text", ""),
                        "computer_read": "" if row.get("added") else cell.get("original", ""),
                        "status": _cell_status(cell, row),
                        "checked_by": names.get(row.get("verified_by") or "", "") if row.get("status") == "VERIFIED" else "",
                        "checked_at": row.get("verified_at") or "" if row.get("status") == "VERIFIED" else "",
                    })
        else:
            out.append({
                "section": {"heading": "Heading", "paragraph": "Text", "image": "Picture"}.get(block["type"], block["type"]),
                "section_type": block["type"],
                "page": _page(block),
                "row": None,
                "field": "",
                "value": block.get("text", ""),
                "computer_read": "" if block.get("added") else block.get("original", block.get("text", "")),
                "status": _block_status(block),
                "checked_by": names.get(block.get("verified_by") or "", "") if block.get("status") == "VERIFIED" else "",
                "checked_at": block.get("verified_at") or "" if block.get("status") == "VERIFIED" else "",
            })
    return out


_COLUMNS = [("section", "Section"), ("page", "Page"), ("row", "Row"), ("field", "Field"), ("value", "Value"),
            ("status", "Status"), ("computer_read", "What the computer read"),
            ("checked_by", "Checked by"), ("checked_at", "Checked at")]


def to_csv(records: List[Dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title for _, title in _COLUMNS])
    for r in records:
        w.writerow(["" if r[k] is None else r[k] for k, _ in _COLUMNS])
    # BOM so Excel opens Devanagari text correctly
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_json(review_doc: Dict[str, Any], records: List[Dict[str, Any]], meta: Dict[str, Any]) -> bytes:
    return json.dumps({**meta, "values": records}, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def to_xlsx(review_doc: Dict[str, Any], records: List[Dict[str, Any]], meta: Dict[str, Any], names: Dict[str, str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    fills = {k: PatternFill("solid", fgColor=v) for k, v in _FILL.items() if v}

    about = wb.active
    about.title = "About"
    counts: Dict[str, int] = {}
    for r in records:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    rows = [
        ("Document", meta["document"]),
        ("Exported", meta["exported_at"]),
        ("Exported by", meta["exported_by"]),
        ("", ""),
        ("Status of each value", "Count"),
    ] + [(s, counts.get(s, 0)) for s in (CHECKED, CORRECTED, ADDED, MACHINE)] + [
        ("", ""),
        ("This is the corrected version. Removed rows and sections are not included.", ""),
        ("Each value has a status column next to it; cells are coloured by status.", ""),
    ]
    for r in rows:
        about.append(list(r))
    for s, fill in fills.items():
        for row in about.iter_rows(min_row=6, max_row=9):
            if row[0].value == s:
                row[0].fill = fill
    about["A5"].font = bold
    about["B5"].font = bold
    about.column_dimensions["A"].width = 60
    about.column_dimensions["B"].width = 40

    used_titles = {"About", "Text", "All values"}
    table_no = 0
    for block in review_doc["blocks"]:
        if block.get("deleted") or block["type"] != "table":
            continue
        table_no += 1
        title = re.sub(r"[\[\]\*\?/\\:]", " ", block.get("title") or f"Table {table_no}")[:28]
        base, k = title, 2
        while title in used_titles:
            title = f"{base[:25]} {k}"
            k += 1
        used_titles.add(title)
        ws = wb.create_sheet(title)
        headers = [_label(h) for h in block.get("headers") or []]
        head = ["Row", "Page", "Row status"]
        for h in headers:
            head += [h, f"{h} - status"]
        ws.append(head)
        for c in ws[1]:
            c.font = bold
            c.alignment = wrap
        live = [r for r in block.get("rows") or [] if not r.get("deleted")]
        for n, row in enumerate(live, start=1):
            statuses = [_cell_status(c, row) for c in row["cells"]]
            row_status = CHECKED if row.get("status") == "VERIFIED" else ADDED if row.get("added") else \
                CORRECTED if CORRECTED in statuses else MACHINE
            line: List[Any] = [n, _page(block, row), row_status]
            for c, st in zip(row["cells"], statuses):
                line += [c.get("text", ""), st]
            ws.append(line)
            excel_row = ws.max_row
            for i, st in enumerate(statuses):
                if st in fills:
                    ws.cell(row=excel_row, column=4 + 2 * i).fill = fills[st]
            if row_status in fills:
                ws.cell(row=excel_row, column=3).fill = fills[row_status]
        for i in range(1, len(head) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 10 if i <= 2 else 22 if i % 2 else 30
        ws.freeze_panes = "D2"

    text_records = [r for r in records if r["section_type"] != "table"]
    if text_records:
        ws = wb.create_sheet("Text")
        ws.append(["Page", "Kind", "Text", "Status", "What the computer read", "Checked by"])
        for c in ws[1]:
            c.font = bold
        for r in text_records:
            ws.append([r["page"], r["section"], r["value"], r["status"], r["computer_read"], r["checked_by"]])
            if r["status"] in fills:
                ws.cell(row=ws.max_row, column=3).fill = fills[r["status"]]
            ws.cell(row=ws.max_row, column=3).alignment = wrap
        ws.column_dimensions["C"].width = 80
        ws.column_dimensions["E"].width = 50

    ws = wb.create_sheet("All values")
    ws.append([title for _, title in _COLUMNS])
    for c in ws[1]:
        c.font = bold
    for r in records:
        ws.append(["" if r[k] is None else r[k] for k, _ in _COLUMNS])
        if r["status"] in fills:
            ws.cell(row=ws.max_row, column=5).fill = fills[r["status"]]
    for letter, width in zip("ABCDEFGHI", (18, 7, 6, 26, 40, 24, 40, 20, 22)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def content_disposition(title: str, fmt: str) -> str:
    stem = re.sub(r"\.(pdf|png|jpe?g|tiff?|bmp|webp)$", "", title or "document", flags=re.I)
    name = f"{stem} - corrected.{fmt}"
    ascii_name = re.sub(r"[^A-Za-z0-9 ._-]", "_", name)
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(name)}"


MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
}


async def export_document(db: AsyncSession, tenant_id: UUID, document_id: UUID, actor_id: UUID, role: str, fmt: str) -> Tuple[bytes, str, str]:
    """Returns (bytes, media type, Content-Disposition). Audited."""
    from fastapi import HTTPException

    from app.models.user import User
    from app.services import review_service
    from app.services.audit_service import log_action

    if fmt not in FORMATS:
        raise HTTPException(status_code=400, detail=f"format must be one of: {', '.join(FORMATS)}")
    review_doc = await review_service.get_review_document(db, tenant_id, document_id, role)

    ids = {b.get("verified_by") for b in review_doc["blocks"]} | {
        r.get("verified_by") for b in review_doc["blocks"] for r in b.get("rows") or []
    }
    ids = {i for i in ids if i}
    ids.add(str(actor_id))
    res = await db.execute(select(User.id, User.full_name, User.email).where(User.id.in_([UUID(i) for i in ids])))
    names = {str(i): (n or e) for i, n, e in res.all()}

    records = flatten(review_doc, names)
    meta = {
        "document": review_doc["title"],
        "document_id": review_doc["document_id"],
        "exported_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"),
        "exported_by": names.get(str(actor_id), str(actor_id)),
        "review_version": review_doc["version"],
        "corrections": review_doc["edit_count"],
    }
    if fmt == "csv":
        body = to_csv(records)
    elif fmt == "json":
        body = to_json(review_doc, records, meta)
    else:
        body = to_xlsx(review_doc, records, meta, names)

    await log_action(
        db, actor_id, tenant_id, "review.export",
        resource_type="document", resource_id=document_id,
        details={"format": fmt, "values": len(records), "review_version": review_doc["version"]},
    )
    return body, MEDIA_TYPES[fmt], content_disposition(review_doc["title"], fmt)
